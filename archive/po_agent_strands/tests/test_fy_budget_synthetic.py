"""The FY budget converter on synthetic workbooks — always runs, no financial data.

`tests/test_fy_budget_conversion.py` reproduces Finance's real FY26 and FY27 runs,
but its fixtures are real budget rows that are never committed, so it skips on a
fresh clone. This file pins the same behaviour on made-up workbooks with the same
layouts (`tests/fy_synthetic.py`). Each expected amount is computed from the
builder's formula, not read back from the converter. The rebuild guide uses this
file as the converter's spec.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from pmagent.tools.fy_budget import convert_budget as cb
from pmagent.tools.fy_budget.detect import detect_roles, detect_roles_with_warnings
from pmagent.tools.fy_budget.fy27_merged import is_fy27_merged_workbook
from pmagent.tools.fy_budget.pipeline import convert_fy_budget
from tests import fy_synthetic as fx

DAYS_IN_FY = 365          # Feb 2025–Jan 2026 and Jul 2026–Jun 2027 have no 29 Feb
ENTITIES = {
    "RQF | Banner State": len(fx.BANNERS),
    "RQF | Customer Group State": len(fx.GROUPS) + len(fx.PLANTS),
    "RQF INCLD CWH | Customer Group State": len(fx.PRIVATE_MG1),
    "RQF INCLD CWH | MG1": len(fx.NONPRIVATE_MG1),
}


@pytest.fixture(scope="module")
def legacy(tmp_path_factory):
    folder = tmp_path_factory.mktemp("legacy")
    core, custplant = fx.write_legacy(folder)
    return core, custplant, folder


@pytest.fixture(scope="module")
def legacy_run(legacy, tmp_path_factory):
    core, custplant, _ = legacy
    out = tmp_path_factory.mktemp("out") / "fy26.csv"
    result = convert_fy_budget([custplant, core], fx.LEGACY_START, out)
    return pd.read_csv(out, keep_default_na=False, low_memory=False), result


@pytest.fixture(scope="module")
def fy27(tmp_path_factory):
    return fx.write_fy27(tmp_path_factory.mktemp("fy27"))


@pytest.fixture(scope="module")
def fy27_run(fy27, tmp_path_factory):
    out = tmp_path_factory.mktemp("out27") / "fy27.csv"
    result = convert_fy_budget([fy27], fx.FY27_START, out)
    return pd.read_csv(out, keep_default_na=False, low_memory=False), result


def monthly(df: pd.DataFrame, mask, key: str) -> pd.DataFrame:
    sub = df[mask].copy()
    sub["ym"] = sub["Cal Year Mth Day"].astype(int) // 100
    return sub.groupby([key, "ym"])[["Budget", "GP"]].sum()


# --- the fiscal calendar -----------------------------------------------------

def test_fy_months_cross_the_year_end():
    assert cb.fy_months(202507) == [202507, 202508, 202509, 202510, 202511, 202512,
                                    202601, 202602, 202603, 202604, 202605, 202606]
    assert len(cb.days_of(202402)) == 29 and cb.days_of(202502)[-1] == 20250228


# --- role detection ----------------------------------------------------------

def test_roles_are_detected_by_content_in_any_order(legacy, fy27):
    core, custplant, _ = legacy
    roles = detect_roles([custplant, core])
    assert roles == {"core": core, "custplant": custplant}
    assert detect_roles([fy27]) == {"fy27_merged": fy27}
    assert is_fy27_merged_workbook(fy27) and not is_fy27_merged_workbook(core)


def test_a_duplicate_role_is_a_warning_not_a_failure(legacy, tmp_path):
    core, custplant, _ = legacy
    copy = tmp_path / "core_copy.xlsx"
    copy.write_bytes(core.read_bytes())
    roles, warnings = detect_roles_with_warnings([core, copy, custplant])
    assert roles["core"] == core
    assert any("Multiple files matched role 'core'" in w for w in warnings)


def test_missing_inputs_name_what_was_seen(legacy):
    core, _, _ = legacy
    with pytest.raises(RuntimeError, match=r"Could not identify required file\(s\): \['custplant'\]"):
        detect_roles([core])


def test_a_daily_rate_table_is_recognised_by_its_columns(tmp_path):
    daily = tmp_path / "weights.csv"
    pd.DataFrame({"Cal Year Mth Day": [20250201], "Daily Weight": [1.0]}).to_csv(daily, index=False)
    other = tmp_path / "notes.csv"
    pd.DataFrame({"a": [1]}).to_csv(other, index=False)
    roles, _ = detect_roles_with_warnings([daily, other, *fx.write_legacy(tmp_path)])
    assert roles["dailyrate"] == daily


# --- the legacy two-file layout ----------------------------------------------

def test_one_row_per_entity_per_day(legacy_run):
    df, result = legacy_run
    assert result["summary"]["rows_by_partition"] == {k: n * DAYS_IN_FY for k, n in ENTITIES.items()}
    assert len(df) == sum(ENTITIES.values()) * DAYS_IN_FY
    assert list(df.columns) == cb.OUTPUT_COLS
    assert df["Cal Year Mth Day"].min() == 20250201 and df["Cal Year Mth Day"].max() == 20260131


def test_internal_validation_passes_and_metadata_is_written(legacy_run):
    _, result = legacy_run
    assert result["validation_failures"] == []
    assert result["summary"]["internal_validation"] == "passed"
    metadata = json.loads(Path(result["metadata_path"]).read_text(encoding="utf-8"))
    assert Path(result["metadata_path"]).name == "fy26.metadata.json"
    assert metadata["schema_validation"] == "passed" and metadata["internal_validation"] == "passed"
    assert metadata["output_schema"] == cb.OUTPUT_COLS
    assert metadata["total_rows"] == sum(ENTITIES.values()) * DAYS_IN_FY
    assert metadata["inputs"]["dailyrate"] is None


@pytest.mark.parametrize("kind,mask,key,codes", [
    ("banner", lambda d: d.Partition == "Banner State", "Banner Group Code", fx.BANNERS),
    ("group", lambda d: (d.Type == "RQF") & (d.Partition == "Customer Group State")
     & (d["Customer Group Code"] != "PLANT"), "Customer Group Code", fx.GROUPS),
    ("plant", lambda d: d["Customer Group Code"] == "PLANT", "Plant Code", tuple(fx.PLANTS)),
    ("private", lambda d: (d.Type == "RQF INCLD CWH") & (d.Partition == "Customer Group State"),
     "MG1 Code", fx.PRIVATE_MG1),
    ("nonprivate", lambda d: d.Partition == "MG1", "MG1 Code", fx.NONPRIVATE_MG1),
])
def test_every_month_rolls_up_to_its_source(legacy_run, kind, mask, key, codes):
    df, _ = legacy_run
    got = monthly(df, mask(df), key)
    assert sorted(got.index.get_level_values(0).unique()) == sorted(codes)
    for code in codes:
        for i, ym in enumerate(fx.months(fx.LEGACY_START)):
            assert got.loc[(code, ym), "Budget"] == pytest.approx(fx.legacy_expected(kind, code, i, "ns"), abs=0.2)
            assert got.loc[(code, ym), "GP"] == pytest.approx(fx.legacy_expected(kind, code, i, "gp"), abs=0.2)


def test_the_cuts_are_stamped_as_the_view_expects(legacy_run):
    df, _ = legacy_run
    private = df[(df.Type == "RQF INCLD CWH") & (df.Partition == "Customer Group State")]
    assert set(private["Private Label Flag"]) == {"Yes"} and set(private["Customer Group Code"]) == {""}
    banner = df[df.Partition == "Banner State"]
    assert set(banner["MG1 Code"]) == {"FOS"} and set(banner["Plant Code"]) == {"NA"}
    assert set(df["Division Code"]) == {2} and set(df["Sales Organisation"]) == {1010}
    assert (df["RQF"] == df["Budget"]).all()
    assert (df["Div_Dist_Key"].astype(str) == "201").all()
    group = df[df["Customer Group Code"] == "GRA"].iloc[0]
    assert group["Group_State_Key"] == "GRANA" and group["MG1_PL_SO_DC_Key"] == "FOSNo011010"


def test_without_a_daily_file_each_month_is_split_evenly(legacy_run):
    df, result = legacy_run
    assert result["summary"]["daily_distribution"] == "even split (no daily file)"
    feb = df[(df["Banner Group Code"] == "AMC") & (df["Cal Year Mth Day"] // 100 == 202502)]
    assert len(feb) == 28 and feb["Budget"].nunique() == 1
    assert feb["Budget"].iloc[0] == round(fx.value("AMC", 0, "ns") / 28, 2)


def _daily_file(folder: Path, weight) -> Path:
    days = [d for ym in fx.months(fx.LEGACY_START) for d in fx.days(ym)]
    path = folder / "daily.csv"
    pd.DataFrame({"date": days, "prec": [weight(d) for d in days]}).to_csv(path, index=False)
    return path


def test_a_full_daily_file_is_applied_and_still_ties_back(legacy, tmp_path):
    core, custplant, _ = legacy
    daily = _daily_file(tmp_path, lambda d: d % 100)            # later days weigh more
    result = convert_fy_budget([core, custplant, daily], fx.LEGACY_START, tmp_path / "w.csv")
    assert result["summary"]["daily_distribution"] == "applied"
    assert result["validation_failures"] == []
    df = pd.read_csv(tmp_path / "w.csv", keep_default_na=False)
    amc = df[df["Banner Group Code"] == "AMC"].set_index("Cal Year Mth Day")["Budget"]
    assert amc[20250228] > amc[20250201]


def test_a_daily_file_that_misses_days_aborts(legacy, tmp_path):
    core, custplant, _ = legacy
    path = tmp_path / "short.csv"
    pd.DataFrame({"date": fx.days(202502), "prec": 1.0}).to_csv(path, index=False)
    with pytest.raises(RuntimeError, match="does not cover the fiscal year"):
        convert_fy_budget([core, custplant, path], fx.LEGACY_START, tmp_path / "x.csv")


def test_a_month_whose_weights_sum_to_zero_aborts(legacy, tmp_path):
    core, custplant, _ = legacy
    daily = _daily_file(tmp_path, lambda d: 0)
    with pytest.raises(RuntimeError, match="weights sum to zero or less"):
        convert_fy_budget([core, custplant, daily], fx.LEGACY_START, tmp_path / "z.csv")


def test_schema_drift_fails_loudly():
    with pytest.raises(RuntimeError, match="Published output schema mismatch"):
        cb.validate_output_schema(pd.DataFrame(columns=cb.OUTPUT_COLS[:-1] + ["Unexpected"]))
    with pytest.raises(RuntimeError, match="Internal output schema mismatch"):
        cb.validate_raw_output_schema(pd.DataFrame(columns=cb.RAW_OUTPUT_COLS[1:]))


# --- the FY27 merged workbook ------------------------------------------------

def test_fy27_converts_with_its_embedded_daily_allocations(fy27_run):
    df, result = fy27_run
    assert result["validation_failures"] == []
    assert result["summary"]["daily_distribution"] == "embedded Core/CW daily allocations"
    assert result["summary"]["rows_by_partition"] == {k: n * DAYS_IN_FY for k, n in ENTITIES.items()}
    assert list(df.columns) == cb.OUTPUT_COLS
    assert df["Cal Year Mth Day"].min() == 20260701 and df["Cal Year Mth Day"].max() == 20270630
    metadata = json.loads(Path(result["metadata_path"]).read_text(encoding="utf-8"))
    assert "business-confirmed rule" in metadata["source_layout"]["customer_daily_allocation"]


def test_fy27_customer_view_uses_the_core_allocation_never_a_blend(fy27_run):
    df, _ = fy27_run
    day = 20260703
    row = df[(df["Banner Group Code"] == "AMC") & (df["Cal Year Mth Day"] == day)]
    assert row["Budget"].iloc[0] == pytest.approx(
        fx.value("AMC", 0, "ns") * fx.fy27_daily_share(day, "core"), abs=0.006)


def test_fy27_non_private_mg1_sums_core_and_cw_each_on_its_own_allocation(fy27_run):
    df, _ = fy27_run
    day, mg1, section = 20260705, "FOS", "Exclude Private Label"
    expected = sum(fx.fy27_category(mg1, section, part, 0, "ns") * fx.fy27_daily_share(day, part)
                   for part in ("core", "cw"))
    row = df[(df.Partition == "MG1") & (df["MG1 Code"] == mg1) & (df["Cal Year Mth Day"] == day)]
    assert row["Budget"].iloc[0] == pytest.approx(expected, abs=0.006)


def test_fy27_months_roll_up_to_the_total_category_view(fy27_run):
    df, _ = fy27_run
    got = monthly(df, df.Partition == "MG1", "MG1 Code")
    for mg1 in fx.NONPRIVATE_MG1:
        for i, ym in enumerate(fx.months(fx.FY27_START)):
            want = fx.fy27_category(mg1, "Exclude Private Label", "total", i, "ns")
            assert got.loc[(mg1, ym), "Budget"] == pytest.approx(want, abs=0.2)


def test_fy27_published_contract(fy27_run):
    df, _ = fy27_run
    assert set(df["State"]) == {"NA"}
    assert set(df["Plant Code"]) == {"NA", *fx.PLANTS}
    assert (df["Ban_State_Key"] == df["Banner Group Code"] + df["State"]).all()


def test_fy27_rejects_a_fiscal_start_that_does_not_match_its_headers(fy27, tmp_path):
    with pytest.raises(RuntimeError, match="month headers are"):
        convert_fy_budget([fy27], 202608, tmp_path / "bad.csv")


def test_fy27_rejects_daily_weights_that_do_not_sum_to_one(fy27, tmp_path):
    wb = load_workbook(fy27)
    wb["FY27 Daily Allocation_CW"]["B2"] = 0.9
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    with pytest.raises(RuntimeError, match="daily weights for 202607 sum to"):
        convert_fy_budget([broken], fx.FY27_START, tmp_path / "b.csv")


# --- the published contract, and validation that can actually fail -----------

def test_published_columns_are_the_view_contract():
    # The Snowflake view ingests exactly these names in exactly this order
    # (pmagent/skills/fy_budget/SKILL.md, "Output contract").
    assert cb.OUTPUT_COLS == [
        "Type", "Partition", "Cal Year Mth Day", "Division Code", "Distr Channel Code",
        "Order Type", "Plant Code", "MG1 Code", "Product Hierarchy Level 5",
        "Sales Organisation", "Customer Group Code", "Private Label Flag",
        "Banner Group Code", "State", "Budget", "GP", "RQF", "Comments",
        "Division Key", "Dist Channel Key", "Div_Dist_Key", "MG1_PL_SO_DC_Key",
        "Ban_State_Key", "Group_State_Key"]


def test_the_derived_keys(legacy_run):
    df, _ = legacy_run
    private = df[(df.Type == "RQF INCLD CWH") & (df.Partition == "Customer Group State")]
    assert set(private["Group_State_Key"]) == {""}                 # no group: the key stays blank
    assert set(private["MG1_PL_SO_DC_Key"]) == {f"{m}Yes011010" for m in fx.PRIVATE_MG1}
    rqf = df[df.Type == "RQF"]
    assert set(rqf["Private Label Flag"]) == {"No"}
    plant = df[df["Customer Group Code"] == "PLANT"].iloc[0]
    assert plant["Group_State_Key"] == "PLANTNA" and plant["Ban_State_Key"] == "NANA"


def _drop_a_day(monkeypatch):
    """Make the daily expansion lose one day, so no entity ties back to its source."""
    real = cb._expand_to_days
    monkeypatch.setattr(cb, "_expand_to_days", lambda *a, **k: real(*a, **k).iloc[:-1])


def test_a_legacy_run_that_does_not_tie_back_is_reported_as_failed(legacy, tmp_path, monkeypatch):
    core, custplant, _ = legacy
    _drop_a_day(monkeypatch)
    result = convert_fy_budget([core, custplant], fx.LEGACY_START, tmp_path / "bad.csv")
    assert result["validation_failures"]
    assert any("Banner AMC 202601" in f for f in result["validation_failures"])
    assert result["summary"]["internal_validation"] == f"{len(result['validation_failures'])} failures"
    metadata = json.loads(Path(result["metadata_path"]).read_text(encoding="utf-8"))
    assert metadata["internal_validation"] != "passed" and metadata["validation_failures"]


def test_an_fy27_run_that_does_not_tie_back_is_refused(fy27, tmp_path, monkeypatch):
    _drop_a_day(monkeypatch)
    with pytest.raises(RuntimeError, match="FY27 internal validation failed"):
        convert_fy_budget([fy27], fx.FY27_START, tmp_path / "bad.csv")


def test_fy27_refuses_a_total_category_view_that_disagrees_with_its_parts(fy27, tmp_path):
    wb = load_workbook(fy27)
    ws = wb["Total Cat View"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "FOS")
    ws.cell(row, 2).value = ws.cell(row, 2).value + 1000
    broken = tmp_path / "disagrees.xlsx"
    wb.save(broken)
    with pytest.raises(RuntimeError, match="Non-private MG1 FOS 202607 does not roll up"):
        convert_fy_budget([broken], fx.FY27_START, tmp_path / "d.csv")


@pytest.mark.parametrize("label,mask,parts", [
    ("plant", lambda d: d["Plant Code"] == "1201", {"core": ("1201", "core"), "cw": ("1201", "cw")}),
    ("customer group", lambda d: d["Customer Group Code"] == "GRA", {"core": ("GRA", "")}),
    ("private MG1", lambda d: (d.Partition == "Customer Group State") & (d["MG1 Code"] == "PBS"),
     {"core": ("PBS", "Private Labelcore")}),
])
def test_fy27_each_component_is_spread_on_its_own_allocation(fy27_run, label, mask, parts):
    df, _ = fy27_run
    day = 20260709
    expected = sum(fx.value(code, 0, "ns", part) * fx.fy27_daily_share(day, allocation)
                   for allocation, (code, part) in parts.items())
    row = df[mask(df) & (df["Cal Year Mth Day"] == day)]
    assert row["Budget"].iloc[0] == pytest.approx(expected, abs=0.006), label


def test_the_pipeline_result_has_the_keys_the_finance_report_reads(fy27, tmp_path):
    # D5's render_run_report reads exactly these (test_finance_tools.py checks the report).
    result = convert_fy_budget([fy27], fx.FY27_START, tmp_path / "ok.csv")
    assert set(result) == {"steps", "summary", "warnings", "validation_failures", "out_path",
                           "metadata_path"}
    assert set(result["summary"]) == {"fiscal_year_start", "fiscal_months", "total_rows",
                                      "rows_by_partition", "daily_distribution", "warnings",
                                      "metadata_path", "internal_validation"}
    assert result["summary"]["fiscal_year_start"] == fx.FY27_START

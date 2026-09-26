"""
Identify which uploaded file is which, by inspecting contents rather than filename.

Supported formats:
  fy27_merged : one workbook containing all FY27 Customer/Plant/Category views
                and embedded Core/CW daily-allocation sheets.
  core      : workbook containing a 'Total by month' sheet  (the MG1 cuts come from here)
  custplant : workbook containing 'Customer view' and 'Plant_View' sheets
  dailyrate : OPTIONAL table with a date column and a weight column (xlsx/csv/tsv)

Returns a dict {role: Path}. A `fy27_merged` workbook is sufficient by itself;
otherwise `core` and `custplant` are required and `dailyrate` is optional.
"""
from __future__ import annotations
import logging
from pathlib import Path

from openpyxl import load_workbook

from . import convert_budget as cb
from .fy27_merged import is_fy27_merged_workbook

logger = logging.getLogger(__name__)


def _sheets(path: Path) -> list[str]:
    try:
        return load_workbook(path, read_only=True, data_only=True).sheetnames
    except Exception:
        return []


def _looks_like_dailyrate(path: Path) -> bool:
    """A small table whose columns include a recognised date alias and weight alias."""
    try:
        import pandas as pd
        suffix = path.suffix.lower()
        if suffix == '.xlsx':
            cols = pd.read_excel(path, nrows=1).columns
        elif suffix in ('.csv', '.tsv', '.txt'):
            sep = '\t' if suffix in ('.tsv', '.txt') else ','
            cols = pd.read_csv(path, nrows=1, sep=sep).columns
        else:
            return False
    except Exception:
        return False
    lower = {str(c).strip().lower() for c in cols}
    has_date = any(a.lower() in lower for a in cb.DR.date_aliases)
    has_weight = any(a.lower() in lower for a in cb.DR.weight_aliases)
    return has_date and has_weight


def detect_roles_with_warnings(paths: list[Path]) -> tuple[dict[str, Path], list[str]]:
    """Detect roles and report duplicate matches without breaking the run."""
    candidates: dict[str, list[Path]] = {'core': [], 'custplant': [], 'dailyrate': [], 'fy27_merged': []}
    for p in paths:
        if is_fy27_merged_workbook(p):
            candidates['fy27_merged'].append(p)
            continue
        sheets = set(_sheets(p))
        if cb.CORE.sheet in sheets:
            candidates['core'].append(p)
        elif {cb.CV.sheet, cb.PV.sheet} <= sheets:
            candidates['custplant'].append(p)
        elif _looks_like_dailyrate(p):
            candidates['dailyrate'].append(p)

    roles: dict[str, Path] = {}
    warnings: list[str] = []
    for role, matches in candidates.items():
        if not matches:
            continue
        roles[role] = matches[0]
        if len(matches) > 1:
            ignored = ', '.join(p.name for p in matches[1:])
            msg = f"Multiple files matched role {role!r}; using {matches[0].name}, ignoring {ignored}."
            warnings.append(msg)
            logger.warning(msg)

    if 'fy27_merged' in roles:
        return roles, warnings
    missing = [r for r in ('core', 'custplant') if r not in roles]
    if missing:
        detail = '; '.join(f"{p.name}: sheets={_sheets(p) or 'n/a'}" for p in paths)
        raise RuntimeError(
            f"Could not identify required file(s): {missing}. "
            f"Need a workbook with a {cb.CORE.sheet!r} sheet (core) and one with "
            f"{cb.CV.sheet!r}+{cb.PV.sheet!r} sheets (custplant). Saw -> {detail}")
    return roles, warnings


def detect_roles(paths: list[Path]) -> dict[str, Path]:
    """Backward-compatible role detection API."""
    roles, _ = detect_roles_with_warnings(paths)
    return roles
